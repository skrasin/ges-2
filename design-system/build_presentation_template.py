#!/usr/bin/env python3
"""
Собирает design-system/ges2-presentation-template.pptx — шаблон презентации
на основе дизайн-системы ГЭС-2, вёрстка проверена на двух реальных колодах
ГЭС-2/V–A–C (см. design-system.md, §9 «Презентации»).

Шрифт: проприетарная "Diagramatika" недоступна — используется Arial
(универсальный кросс-платформенный fallback, тот же принцип, что и в
остальных Office-документах проекта).

Использование:
    python3 design-system/build_presentation_template.py
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
PHOTOS = HERE / "photos"
DIAGRAMS = HERE / "diagrams"
OUT = HERE / "ges2-presentation-template.pptx"

FONT = "Arial"
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x68, 0x62)
DIAGRAM_BLUE = RGBColor(0x00, 0x75, 0xB9)
DIAGRAM_RED = RGBColor(0xB4, 0x17, 0x00)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)


def set_tracking(run, plus_percent=0.02):
    """Diagramatika Display несёт трекинг +2% от кегля — воспроизводим тем же
    правилом на fallback-шрифте. Величина в сотых долях пункта (атрибут spc)."""
    size_pt = run.font.size.pt if run.font.size else 18
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(round(size_pt * plus_percent * 100))))


def add_text(slide, left, top, width, height, text, size, bold=False,
             color=BLACK, align=PP_ALIGN.LEFT, display=False, anchor=None,
             line_spacing=1.15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        if display:
            set_tracking(r)
    return box


def add_hairline(slide, x1, y, x2, color=BLACK, weight=0.75):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y, x2, y)
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    return conn


def add_breadcrumb(slide, section=""):
    text = "V—A—C > ГЭС-2" + (f" > {section}" if section else "")
    add_text(slide, MARGIN, Inches(0.35), Inches(9), Inches(0.35),
              text, size=11, color=MUTED)


def add_page_number(slide, n):
    add_text(slide, SLIDE_W - Inches(1.6), Inches(0.35), Inches(1.0), Inches(0.35),
              f"стр. {n}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def add_headline(slide, text, top=Inches(0.75), size=40, width=None):
    add_text(slide, MARGIN, top, width or (SLIDE_W - 2 * MARGIN), Inches(1.6),
              text, size=size, display=True)


def cover_crop_box(slide, image_path, left, top, width, height):
    """Вставляет фото с заполнением всей рамки без искажений пропорций
    (аналог CSS object-fit: cover) — обрезка лишнего по краям."""
    with Image.open(image_path) as im:
        img_w, img_h = im.size
    img_ar = img_w / img_h
    box_ar = width / height

    pic = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)

    if img_ar > box_ar:
        kept_w = img_h * box_ar
        crop = max(0.0, (1 - kept_w / img_w) / 2)
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        kept_h = img_w / box_ar
        crop = max(0.0, (1 - kept_h / img_h) / 2)
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)  # MSO_SHAPE.RECTANGLE
    rect.fill.solid()
    rect.fill.fore_color.rgb = WHITE
    rect.line.fill.background()
    rect.shadow.inherit = False
    # переносим фон на задний план
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 0. Инструкция по использованию шаблона
    s = blank_slide(prs)
    add_breadcrumb(s, "Шаблон презентации")
    add_headline(s, "Как пользоваться этим шаблоном", size=36)
    instructions = (
        "Слайд 1 — титул. Слайды 2—6 — примеры вёрстки, встреченные в реальных "
        "презентациях ГЭС-2/V–A–C: текст + фото, буллеты + фото, два фото рядом, "
        "тайминг, план помещения. Слайд 7 — контакты/закрытие.\n\n"
        "Продублируйте нужный слайд (не создавайте с нуля) и замените "
        "текст и фото в квадратных скобках — расположение, кегль, отступы и "
        "hairline-линии менять не нужно, они уже соответствуют дизайн-системе.\n\n"
        "Шрифт: проприетарная Diagramatika недоступна вне сайта ges-2.org — "
        "здесь используется Arial как ближайший универсальный аналог "
        "(тот же fallback, что и в остальных документах проекта).\n\n"
        "Эту инструкцию перед отправкой презентации удалите."
    )
    add_text(s, MARGIN, Inches(1.9), SLIDE_W - 2 * MARGIN, Inches(4.5),
              instructions, size=16, line_spacing=1.3)

    # 1. Титул
    s = blank_slide(prs)
    add_breadcrumb(s)
    add_text(s, MARGIN, Inches(0.7), Inches(6), Inches(0.4),
              "Дом культуры «ГЭС-2»", size=14, color=MUTED)
    add_text(s, MARGIN, Inches(2.6), SLIDE_W - 2 * MARGIN, Inches(2.6),
              "[Название презентации]", size=54, display=True)
    add_text(s, MARGIN, Inches(5.6), Inches(6), Inches(0.5),
              "[Подзаголовок / для кого]", size=18, color=MUTED)
    add_text(s, MARGIN, Inches(6.6), Inches(4), Inches(0.4),
              "[дата]", size=16)

    # 2. Текст + фото
    s = blank_slide(prs)
    add_breadcrumb(s, "О площадке")
    add_page_number(s, 2)
    add_headline(s, "[Заголовок раздела]", size=36)
    photo_left = MARGIN
    photo_top = Inches(1.9)
    photo_w = Inches(6.0)
    photo_h = Inches(5.0)
    cover_crop_box(s, PHOTOS / "exterior-facade-river-chimneys.jpeg",
                    photo_left, photo_top, photo_w, photo_h)
    add_text(s, photo_left + photo_w + Inches(0.5), photo_top,
              SLIDE_W - MARGIN - (photo_left + photo_w + Inches(0.5)), photo_h,
              "[Один-два абзаца текста рядом с фотографией. Фото — на всю "
              "высоту контентной зоны, без рамки и подписи, край в край с "
              "текстовым блоком.]\n\n[Второй абзац при необходимости.]",
              size=16, line_spacing=1.3)

    # 3. Буллеты + фото
    s = blank_slide(prs)
    add_breadcrumb(s, "Программа")
    add_page_number(s, 3)
    add_headline(s, "[Название раздела]", size=36)
    bullets = "\n".join(f"•  [Пункт {i}]" for i in range(1, 6))
    add_text(s, MARGIN, Inches(1.9), Inches(5.6), Inches(5.0),
              bullets, size=18, line_spacing=1.5)
    cover_crop_box(s, PHOTOS / "concert-hall-piano-recital.jpeg",
                    Inches(7.0), Inches(1.9), SLIDE_W - MARGIN - Inches(7.0), Inches(5.0))

    # 4. Два фото рядом
    s = blank_slide(prs)
    add_breadcrumb(s, "Пространства")
    add_page_number(s, 4)
    add_headline(s, "[Локация 1] > [Локация 2]", size=36)
    gap = Inches(0.3)
    pair_w = (SLIDE_W - 2 * MARGIN - gap) / 2
    pair_h = Inches(5.0)
    pair_top = Inches(1.9)
    cover_crop_box(s, PHOTOS / "atrium-cafe-green-crane.jpeg",
                    MARGIN, pair_top, pair_w, pair_h)
    cover_crop_box(s, PHOTOS / "reception-crowd-atrium-artwork.jpeg",
                    MARGIN + pair_w + gap, pair_top, pair_w, pair_h)

    # 5. Тайминг
    s = blank_slide(prs)
    add_breadcrumb(s, "Тайминг")
    add_page_number(s, 5)
    add_headline(s, "Тайминг", size=36)
    blocks = [
        ("[Блок 1]", ["00.00 — 00.00: [что происходит].",
                       "00.00 — 00.00: [что происходит]."]),
        ("[Блок 2]", ["00.00 — 00.00: [что происходит]."]),
        ("[Блок 3]", ["00.00 — 00.00: [что происходит].",
                       "00.00 — [Закрытие]."]),
    ]
    y = Inches(2.1)
    for title, lines in blocks:
        add_text(s, MARGIN, y, Inches(9), Inches(0.5), title, size=18, bold=True)
        y += Inches(0.55)
        for line in lines:
            add_text(s, MARGIN, y, Inches(9), Inches(0.45), line, size=16)
            y += Inches(0.45)
        y += Inches(0.35)

    # 6. Карта / план помещения
    s = blank_slide(prs)
    add_breadcrumb(s, "О площадке")
    add_page_number(s, 6)
    add_headline(s, "Карта Дома культуры «ГЭС-2»", size=32, width=Inches(9))
    add_text(s, SLIDE_W - Inches(2.6), Inches(0.75), Inches(2.0), Inches(0.6),
              "[1 этаж]", size=24, display=True, align=PP_ALIGN.RIGHT)
    plan_top = Inches(1.9)
    plan_h = Inches(4.6)
    plan_w = Inches(8.0)
    with Image.open(DIAGRAMS / "floor-plan-1st-floor.png") as im:
        pw, ph = im.size
    plan_h_final = min(plan_h, plan_w * ph / pw)
    plan_w_final = plan_h_final * pw / ph
    s.shapes.add_picture(str(DIAGRAMS / "floor-plan-1st-floor.png"),
                          MARGIN, plan_top, width=plan_w_final, height=plan_h_final)
    # пример выноски синим — основной маршрут
    line1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(1.2), Inches(2.6), Inches(4.0), Inches(3.0))
    line1.line.color.rgb = DIAGRAM_BLUE
    line1.line.width = Pt(2)
    add_text(s, Inches(0.9), Inches(2.15), Inches(3.5), Inches(0.4),
              "[Основной маршрут]", size=14, color=DIAGRAM_BLUE, bold=True)
    # легенда
    legend_x = SLIDE_W - Inches(3.6)
    legend_y = Inches(2.0)
    for label, color in (("— основной маршрут / зона", DIAGRAM_BLUE),
                          ("— альтернативный / акцент", DIAGRAM_RED)):
        add_text(s, legend_x, legend_y, Inches(3.4), Inches(0.4), label, size=13, color=color)
        legend_y += Inches(0.45)
    add_text(s, legend_x, legend_y + Inches(0.1), Inches(3.4), Inches(0.8),
              "Цвета — только для выносок на планах, не для UI и не бренд-акцент.",
              size=11, color=MUTED, line_spacing=1.3)

    # 7. Контакты / закрытие
    s = blank_slide(prs)
    add_breadcrumb(s, "Контакты")
    add_page_number(s, 7)
    add_headline(s, "Контакты и адрес:", size=36)
    add_text(s, MARGIN, Inches(4.2), Inches(7), Inches(1.2),
              "[Название отдела]\n[email@v-a-c.org]", size=18, line_spacing=1.5)

    prs.save(OUT)
    print(f"OK: {OUT} собран, {len(prs.slides)} слайдов.")


if __name__ == "__main__":
    build()
